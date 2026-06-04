import os
import shutil
import re
import hashlib
import traceback
from langchain_core.embeddings import Embeddings

from typing import List

from fastapi import FastAPI, UploadFile, File, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel

from langchain_google_genai import ChatGoogleGenerativeAI
from langchain_chroma import Chroma
from langchain_core.prompts import ChatPromptTemplate
from langchain_core.runnables import RunnablePassthrough
from langchain_core.output_parsers import StrOutputParser
from langchain_core.documents import Document
from langchain_core.vectorstores import VectorStore

import threading

_init_lock = threading.Lock()

app = FastAPI(title="smartuni RAG API (Powered by Gemini)")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

os.environ["GOOGLE_API_KEY"] = "AIzaSyD64xUsN_XnhPXBDU-ZWyf3c623AUrfWtI"

llm = ChatGoogleGenerativeAI(
    model="gemini-1.5-flash",
    temperature=0,
    google_api_key=os.environ.get("GOOGLE_API_KEY"),
)

class LocalHashEmbeddingFunction(Embeddings):
    """Local embedding using stable feature hashing. No external API calls."""

    def __init__(self, dim: int = 256):
        self.dim = dim

    def _hash_token(self, token: str) -> int:
        h = hashlib.md5(token.encode("utf-8")).hexdigest()
        return int(h, 16) % self.dim

    def embed_query(self, text: str) -> List[float]:
        return self._hash_text(text)

    def embed_documents(self, texts: List[str]) -> List[List[float]]:
        return [self._hash_text(t) for t in texts]

    def _hash_text(self, text: str) -> List[float]:
        words = re.findall(r"[\w\u0600-\u06FF]+", text.lower())
        vec = [0.0] * self.dim
        for w in words:
            idx = self._hash_token(w)
            vec[idx] += 1.0
        total = sum(v for v in vec)
        if total > 0:
            vec = [v / total for v in vec]
        return vec


embedding_fn = LocalHashEmbeddingFunction()

vectorstore: VectorStore | None = None
retriever = None
rag_chain = None

def _init_vectorstore():
    global vectorstore, retriever, rag_chain
    if vectorstore is not None:
        return
    with _init_lock:
        if vectorstore is not None:
            return
        vectorstore = Chroma(persist_directory="./chroma_db", embedding_function=embedding_fn)
        retriever = vectorstore.as_retriever(search_kwargs={"k": 3})

        template = (
            "You are a helpful college teaching assistant. "
            "Answer the student's question based ONLY on the following context. "
            "If the answer is not contained in the context, strictly say "
            "'I'm sorry, that information is not in the uploaded course materials.' "
            "Do not use outside knowledge.\n\n"
            "Context:\n{context}\n\nQuestion: {question}\n\nAnswer:"
        )
        prompt = ChatPromptTemplate.from_template(template)

        rag_chain = (
            {"context": retriever | format_docs, "question": RunnablePassthrough()}
            | prompt
            | llm
            | StrOutputParser()
        )


def get_rag_chain():
    _init_vectorstore()
    if rag_chain is None:
        raise HTTPException(status_code=503, detail="RAG system is initializing. Please try again.")
    return rag_chain

def format_docs(docs):
    return "\n\n".join(doc.page_content for doc in docs)


class ChatRequest(BaseModel):
    question: str


def split_text(text: str, chunk_size: int = 1000, chunk_overlap: int = 100) -> List[str]:
    chunks: List[str] = []
    start = 0
    while start < len(text):
        end = start + chunk_size
        chunk = text[start:end]
        if chunk.strip():
            chunks.append(chunk.strip())
        start = end - chunk_overlap
        if start >= len(text):
            break
    return chunks


def split_documents(documents, chunk_size=1000, chunk_overlap=100):
    out = []
    for doc in documents:
        for piece in split_text(doc.page_content, chunk_size, chunk_overlap):
            out.append(Document(page_content=piece, metadata=doc.metadata))
    return out


@app.get("/")
async def root():
    return {"status": "ok", "docs": "/docs"}


@app.on_event("startup")
def _warm_up():
    try:
        _init_vectorstore()
    except Exception:
        pass


@app.post("/upload")
async def upload_document(file: UploadFile = File(...)):
    filename = (file.filename or "").lower()
    try:
        is_pdf = filename.endswith(".pdf")
        is_docx = filename.endswith(".docx")
        if not (is_pdf or is_docx):
            raise HTTPException(
                status_code=400,
                detail="عفواً، النظام مهيأ حالياً لاستقبال ملفات PDF و Word (.docx) فقط.",
            )

        if vectorstore is None:
            raise HTTPException(status_code=503, detail="Vector store is initializing. Please try again.")

        tmp = f"temp_{file.filename}"
        with open(tmp, "wb") as buf:
            shutil.copyfileobj(file.file, buf)

        try:
            if is_pdf:
                from pypdf import PdfReader
                reader = PdfReader(tmp)
                parts = []
                for i, page in enumerate(reader.pages):
                    txt = page.extract_text()
                    if txt:
                        parts.append(f"\n--- Page {i+1} ---\n{txt}")
                full_text = "".join(parts) if parts else f"[PDF has {len(reader.pages)} pages but no extractable text]"
                documents = [Document(page_content=full_text)]
            else:
                from docx import Document as DocxDocument
                doc = DocxDocument(tmp)
                txt = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
                full_text = txt or "[DOCX file appears empty]"
                documents = [Document(page_content=full_text)]

            chunks = split_documents(documents)
            if not chunks:
                raise HTTPException(
                    status_code=400,
                    detail="عفواً، لم يتم العثور على نصوص قابلة للقراءة في الملف.",
                )

            vectorstore.add_documents(chunks)
            return {
                "message": f"Successfully processed {len(chunks)} chunks from file: {file.filename}."
            }
        finally:
            if os.path.exists(tmp):
                os.remove(tmp)
    except HTTPException:
        raise
    except Exception as e:
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=repr(e))


@app.post("/chat")
async def chat(request: ChatRequest):
    try:
        chain = get_rag_chain()
        answer = chain.invoke(request.question)
        return {"answer": answer}
    except HTTPException:
        raise
    except Exception as e:
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=repr(e))
